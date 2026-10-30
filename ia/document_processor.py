import io
from typing import Tuple, Dict, Any
from pypdf import PdfReader
from docx import Document


class DocumentProcessor:
    """
    Clase encargada de leer archivos PDF y Word (.docx),
    validar su longitud y extraer su contenido textual.
    """

    # Límite de palabras recomendado para no saturar llama3.2:1b en local (~7 a 10 páginas)
    MAX_WORDS_LIMIT: int = 3500

    @classmethod
    def extract_text(cls, file_source, file_name: str) -> Tuple[bool, str, Dict[str, Any]]:
        """
        Extrae el texto de un archivo PDF o Word y aplica el Guardrail de tamaño.

        :param file_source: Ruta del archivo (str) o archivo en memoria (BytesIO/UploadedFile).
        :param file_name: Nombre del archivo con su extensión (.pdf o .docx).
        :return: Tupla (éxito: bool, texto_o_mensaje: str, métricas: dict)
        """
        file_name_lower = file_name.lower()
        text = ""

        try:
            # 1. Extracción según el formato
            if file_name_lower.endswith(".pdf"):
                reader = PdfReader(file_source)
                num_pages = len(reader.pages)
                extracted_pages = []
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        extracted_pages.append(page_text)
                text = "\n".join(extracted_pages)
                metadata = {"tipo": "PDF", "paginas": num_pages}

            elif file_name_lower.endswith(".docx"):
                doc = Document(file_source)
                extracted_chunks = []

                # 1. Extraer párrafos regulares
                for p in doc.paragraphs:
                    cleaned_p = p.text.strip()
                    if cleaned_p:
                        extracted_chunks.append(cleaned_p)

                # 2. Extraer texto dentro de tablas (común en plantillas y ensayos formateados)
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            cell_text = cell.text.strip()
                            if cell_text and cell_text not in extracted_chunks:
                                extracted_chunks.append(cell_text)

                # 3. Extraer cuadros de texto (Text Boxes / Formas de Word)
                try:
                    for txbx in doc.element.xpath('.//w:txbxContent//w:p'):
                        txbx_text = "".join(txbx.itertext()).strip()
                        if txbx_text and txbx_text not in extracted_chunks:
                            extracted_chunks.append(txbx_text)
                except Exception:
                    pass

                text = "\n\n".join(extracted_chunks)
                metadata = {
                    "tipo": "Word (DOCX)",
                    "parrafos": len(doc.paragraphs),
                    "tablas": len(doc.tables)
                }

            else:
                return False, f"Formato no soportado: '{file_name}'. Solo se aceptan archivos .pdf y .docx.", {}

            # Validar que el archivo contenga texto legible
            cleaned_text = text.strip()
            if not cleaned_text:
                return False, "El archivo parece estar vacío o contiene solo imágenes sin texto seleccionable.", {}

            # 2. Conteo de palabras y validación de Guardrail
            words_count = len(cleaned_text.split())
            metadata["palabras"] = words_count

            # Validación de contenido mínimo para evitar alucinaciones con archivos casi vacíos
            if words_count < 15:
                return False, f"El documento contiene muy poco texto ({words_count} palabras). Asegúrate de que el archivo contenga texto editable y no solo capturas o imágenes escaneadas.", metadata

            # 3. Aplicación del Guardrail (Límite de memoria)
            if words_count > cls.MAX_WORDS_LIMIT:
                mensaje_guardrail = (
                    f"😅 ¡Mi memoria no da para tanto!\n\n"
                    f"Este documento contiene **{words_count:,} palabras** "
                    f"(el límite máximo recomendado para mi modelo local es de {cls.MAX_WORDS_LIMIT:,} palabras).\n\n"
                    f"👉 **Solución:** Por favor sube un archivo más corto, un resumen previo, o un capítulo específico."
                )
                return False, mensaje_guardrail, metadata

            # Si pasa la validación, devolvemos el texto listo para la IA
            return True, cleaned_text, metadata

        except Exception as e:
            return False, f"Error al procesar el archivo '{file_name}': {str(e)}", {}
